"""
File Buddy - conversion engine.

Every converter takes (src_path, out_dir, target_ext, opts) and returns a list
of output file paths. The registry below decides which targets each input
type can be turned into.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

# ---------------------------------------------------------------------------
# File type groups
# ---------------------------------------------------------------------------
IMAGE_IN = {"png", "jpg", "jpeg", "webp", "bmp", "gif", "tif", "tiff", "ico", "heic", "heif", "avif"}
IMAGE_OUT = ["png", "jpg", "webp", "bmp", "gif", "tiff", "ico", "pdf"]
PDF_IN = {"pdf"}
PDF_OUT = ["png", "jpg", "webp", "tiff", "txt", "docx", "html"]
OFFICE_IN = {"docx", "doc", "odt", "rtf", "pptx", "ppt", "odp", "xls", "ods"}
TEXT_IN = {"txt", "md"}
TEXT_OUT = ["pdf", "html", "docx", "txt"]
DATA_IN = {"csv", "xlsx", "json"}
DATA_OUT = ["csv", "xlsx", "json", "html"]
AUDIO_IN = {"mp3", "wav", "flac", "ogg", "m4a", "aac", "wma", "opus"}
AUDIO_OUT = ["mp3", "wav", "flac", "ogg", "m4a", "aac"]
VIDEO_IN = {"mp4", "mov", "avi", "mkv", "webm", "wmv", "flv", "m4v"}
VIDEO_OUT = ["mp4", "mov", "mkv", "webm", "avi", "gif", "mp3", "wav"]
MODEL_IN = {"stl", "obj", "ply", "glb", "gltf", "off", "3mf", "dae"}
MODEL_OUT = ["stl", "obj", "ply", "glb", "off", "3mf"]

# Special targets for images (handled by the 3D studio)
IMAGE_3D_TARGETS = ["stl (relief / lithophane)"]


class ConversionError(Exception):
    pass


def ext_of(path: str | Path) -> str:
    return Path(path).suffix.lower().lstrip(".")


def category_of(path: str | Path) -> str:
    e = ext_of(path)
    if e in IMAGE_IN:
        return "image"
    if e in PDF_IN:
        return "pdf"
    if e in OFFICE_IN:
        return "office"
    if e in TEXT_IN:
        return "text"
    if e in DATA_IN:
        return "data"
    if e in AUDIO_IN:
        return "audio"
    if e in VIDEO_IN:
        return "video"
    if e in MODEL_IN:
        return "model"
    return "unknown"


def targets_for(path: str | Path) -> list[str]:
    """List every format this file can be converted to."""
    cat = category_of(path)
    e = ext_of(path)
    e_norm = "jpg" if e == "jpeg" else ("tiff" if e == "tif" else e)
    if cat == "image":
        out = IMAGE_OUT + IMAGE_3D_TARGETS
    elif cat == "pdf":
        out = PDF_OUT
    elif cat == "office":
        out = ["pdf", "txt"] if e in {"docx"} else ["pdf"]
        if e in {"xls", "ods"}:
            out = ["pdf", "xlsx", "csv"]
    elif cat == "text":
        out = TEXT_OUT
    elif cat == "data":
        out = DATA_OUT
    elif cat == "audio":
        out = AUDIO_OUT
    elif cat == "video":
        out = VIDEO_OUT
    elif cat == "model":
        out = MODEL_OUT
    else:
        out = []
    return [t for t in out if t != e_norm]


def unique_path(out_dir: Path, stem: str, ext: str) -> Path:
    p = out_dir / f"{stem}.{ext}"
    i = 1
    while p.exists():
        p = out_dir / f"{stem} ({i}).{ext}"
        i += 1
    return p


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------
def _open_image(src: Path):
    from PIL import Image

    if ext_of(src) in {"heic", "heif", "avif"}:
        try:
            import pillow_heif  # type: ignore

            pillow_heif.register_heif_opener()
        except ImportError as exc:
            raise ConversionError("HEIC/AVIF needs: pip install pillow-heif") from exc
    return Image.open(src)


def convert_image(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    from PIL import Image, ImageSequence

    img = _open_image(src)
    quality = int(opts.get("quality", 92))
    out = unique_path(out_dir, src.stem, target)

    if target == "pdf":
        frames = [f.convert("RGB") for f in ImageSequence.Iterator(img)]
        frames[0].save(out, "PDF", save_all=True, append_images=frames[1:], resolution=300)
        return [out]

    if target in {"jpg", "bmp", "pdf"}:
        # No transparency in these formats -> paste on white
        if img.mode in ("RGBA", "LA", "P"):
            img = img.convert("RGBA")
            bg = Image.new("RGB", img.size, (255, 255, 255))
            bg.paste(img, mask=img.split()[-1])
            img = bg
        else:
            img = img.convert("RGB")

    if target == "jpg":
        img.save(out, "JPEG", quality=quality, optimize=True, subsampling=0)
    elif target == "webp":
        img.save(out, "WEBP", quality=quality, method=6)
    elif target == "ico":
        img.convert("RGBA").save(out, "ICO", sizes=[(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
    elif target == "gif":
        img.save(out, "GIF", save_all=getattr(img, "is_animated", False))
    elif target == "tiff":
        img.save(out, "TIFF", compression="tiff_lzw")
    else:
        img.save(out, target.upper() if target != "jpg" else "JPEG")
    return [out]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def convert_pdf(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    import pymupdf as fitz

    doc = fitz.open(src)
    outs: list[Path] = []
    if target in {"png", "jpg", "webp", "tiff"}:
        dpi = int(opts.get("dpi", 200))
        multi = doc.page_count > 1
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            stem = f"{src.stem}_page{i + 1:03d}" if multi else src.stem
            out = unique_path(out_dir, stem, target)
            if target == "png":
                pix.save(out)
            else:
                from PIL import Image

                im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                if target == "jpg":
                    im.save(out, "JPEG", quality=int(opts.get("quality", 92)))
                else:
                    im.save(out, target.upper())
            outs.append(out)
    elif target == "txt":
        out = unique_path(out_dir, src.stem, "txt")
        text = "\n\n".join(f"----- Page {i + 1} -----\n{p.get_text()}" for i, p in enumerate(doc))
        out.write_text(text, encoding="utf-8")
        outs.append(out)
    elif target == "html":
        out = unique_path(out_dir, src.stem, "html")
        body = "\n".join(p.get_text("html") for p in doc)
        out.write_text(f"<!doctype html><meta charset='utf-8'><title>{src.stem}</title>\n{body}", encoding="utf-8")
        outs.append(out)
    elif target == "docx":
        doc.close()
        from pdf2docx import Converter

        out = unique_path(out_dir, src.stem, "docx")
        cv = Converter(str(src))
        cv.convert(str(out))
        cv.close()
        outs.append(out)
    else:
        raise ConversionError(f"PDF -> {target} not supported")
    return outs


# ---------------------------------------------------------------------------
# Office docs (LibreOffice does the heavy lifting if installed)
# ---------------------------------------------------------------------------
def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    return next((c for c in candidates if os.path.exists(c)), None)


def convert_office(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    if target == "txt" and ext_of(src) == "docx":
        import docx

        d = docx.Document(str(src))
        out = unique_path(out_dir, src.stem, "txt")
        out.write_text("\n".join(p.text for p in d.paragraphs), encoding="utf-8")
        return [out]
    if target in {"xlsx", "csv"} and ext_of(src) in {"xls", "ods"}:
        import pandas as pd

        df = pd.read_excel(src)
        out = unique_path(out_dir, src.stem, target)
        df.to_excel(out, index=False) if target == "xlsx" else df.to_csv(out, index=False)
        return [out]

    soffice = find_soffice()
    if not soffice:
        raise ConversionError(
            "Office files need LibreOffice (free) for this conversion. "
            "Install it from libreoffice.org and try again."
        )
    tmp = out_dir / ".filebuddy_tmp"
    tmp.mkdir(exist_ok=True)
    subprocess.run(
        [soffice, "--headless", "--convert-to", target, "--outdir", str(tmp), str(src)],
        check=True, capture_output=True, timeout=300,
    )
    produced = list(tmp.glob(f"{src.stem}.{target}"))
    if not produced:
        raise ConversionError("LibreOffice didn't produce an output file.")
    out = unique_path(out_dir, src.stem, target)
    shutil.move(str(produced[0]), out)
    shutil.rmtree(tmp, ignore_errors=True)
    return [out]


# ---------------------------------------------------------------------------
# Text / Markdown
# ---------------------------------------------------------------------------
def convert_text(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    raw = src.read_text(encoding="utf-8", errors="replace")
    out = unique_path(out_dir, src.stem, target)
    is_md = ext_of(src) == "md"

    def to_html() -> str:
        if is_md:
            import markdown

            body = markdown.markdown(raw, extensions=["tables", "fenced_code"])
        else:
            import html

            body = f"<pre style='white-space:pre-wrap;font-family:inherit'>{html.escape(raw)}</pre>"
        return (
            "<!doctype html><html><head><meta charset='utf-8'>"
            f"<title>{src.stem}</title><style>body{{font-family:Arial,sans-serif;"
            "max-width:800px;margin:40px auto;line-height:1.5}</style></head>"
            f"<body>{body}</body></html>"
        )

    if target == "html":
        out.write_text(to_html(), encoding="utf-8")
    elif target == "txt":
        out.write_text(raw, encoding="utf-8")
    elif target == "pdf":
        import pymupdf as fitz

        story = fitz.Story(html=to_html())
        writer = fitz.DocumentWriter(str(out))
        rect = fitz.paper_rect("letter")
        where = rect + (54, 54, -54, -54)
        more = True
        while more:
            dev = writer.begin_page(rect)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
        writer.close()
    elif target == "docx":
        import docx

        d = docx.Document()
        for line in raw.splitlines():
            if is_md and line.startswith("#"):
                level = min(len(line) - len(line.lstrip("#")), 4)
                d.add_heading(line.lstrip("#").strip(), level=level)
            else:
                d.add_paragraph(line)
        d.save(str(out))
    return [out]


# ---------------------------------------------------------------------------
# Spreadsheets / data
# ---------------------------------------------------------------------------
def convert_data(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    import pandas as pd

    e = ext_of(src)
    if e == "csv":
        df = pd.read_csv(src)
    elif e == "xlsx":
        df = pd.read_excel(src)
    else:
        df = pd.read_json(src)
    out = unique_path(out_dir, src.stem, target)
    if target == "csv":
        df.to_csv(out, index=False)
    elif target == "xlsx":
        df.to_excel(out, index=False)
    elif target == "json":
        df.to_json(out, orient="records", indent=2)
    elif target == "html":
        df.to_html(out, index=False)
    return [out]


# ---------------------------------------------------------------------------
# Audio / video (ffmpeg comes bundled through imageio-ffmpeg)
# ---------------------------------------------------------------------------
def ffmpeg_exe() -> str:
    p = shutil.which("ffmpeg")
    if p:
        return p
    try:
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception as exc:  # noqa: BLE001
        raise ConversionError("ffmpeg not found. Run: pip install imageio-ffmpeg") from exc


def convert_media(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    out = unique_path(out_dir, src.stem, target)
    cmd = [ffmpeg_exe(), "-y", "-hide_banner", "-loglevel", "error", "-i", str(src)]
    if target == "gif":
        fps = int(opts.get("gif_fps", 12))
        width = int(opts.get("gif_width", 480))
        cmd += ["-vf", f"fps={fps},scale={width}:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse", "-loop", "0"]
    elif target in {"mp3", "wav", "flac", "ogg", "m4a", "aac"}:
        cmd += ["-vn"]
        codec = {"mp3": ["-c:a", "libmp3lame", "-q:a", "2"], "ogg": ["-c:a", "libvorbis", "-q:a", "6"],
                 "m4a": ["-c:a", "aac", "-b:a", "256k"], "aac": ["-c:a", "aac", "-b:a", "256k"],
                 "flac": ["-c:a", "flac"], "wav": ["-c:a", "pcm_s16le"]}[target]
        cmd += codec
    elif target == "webm":
        cmd += ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-c:a", "libopus"]
    elif target in {"mp4", "mov", "mkv"}:
        cmd += ["-c:v", "libx264", "-crf", "20", "-preset", "medium", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k"]
        if target == "mp4":
            cmd += ["-movflags", "+faststart"]
    elif target == "avi":
        cmd += ["-c:v", "mpeg4", "-q:v", "3", "-c:a", "libmp3lame"]
    cmd.append(str(out))
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        raise ConversionError(r.stderr.strip()[-500:] or "ffmpeg failed")
    return [out]


# ---------------------------------------------------------------------------
# 3D models
# ---------------------------------------------------------------------------
def load_mesh(src: Path):
    import trimesh

    m = trimesh.load(str(src), force="scene")
    if isinstance(m, trimesh.Scene):
        if not m.geometry:
            raise ConversionError("That 3D file has no geometry in it.")
        return m
    return m


def convert_model(src: Path, out_dir: Path, target: str, opts: dict) -> list[Path]:
    import trimesh

    scene = load_mesh(src)
    out = unique_path(out_dir, src.stem, target)
    if target in {"stl", "off", "ply", "3mf"}:
        # These want one solid mesh
        mesh = scene.to_geometry() if isinstance(scene, trimesh.Scene) else scene
        mesh.export(str(out), file_type=target)
    else:
        scene.export(str(out), file_type=target)
    return [out]


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------
def convert(src: str | Path, out_dir: str | Path, target: str, opts: dict | None = None) -> list[Path]:
    src, out_dir = Path(src), Path(out_dir)
    opts = opts or {}
    if not src.is_file():
        raise ConversionError(
            f'"{src.name}" is not available at the dropped location. '
            "If it is in OneDrive, wait for it to finish syncing, then try again."
        )
    out_dir.mkdir(parents=True, exist_ok=True)
    cat = category_of(src)

    if target.startswith("stl (relief"):
        from .image3d import image_to_relief_stl

        return [image_to_relief_stl(src, out_dir, opts)]

    fn = {
        "image": convert_image, "pdf": convert_pdf, "office": convert_office,
        "text": convert_text, "data": convert_data, "audio": convert_media,
        "video": convert_media, "model": convert_model,
    }.get(cat)
    if fn is None:
        raise ConversionError(f"File Buddy doesn't know .{ext_of(src)} files yet.")
    if target not in targets_for(src):
        raise ConversionError(f".{ext_of(src)} can't be turned into .{target}")
    return fn(src, out_dir, target, opts)


if __name__ == "__main__":  # tiny CLI: python -m filebuddy.converters in.pdf png
    if len(sys.argv) < 3:
        print("usage: python -m filebuddy.converters <file> <target> [out_dir]")
        sys.exit(1)
    res = convert(sys.argv[1], sys.argv[3] if len(sys.argv) > 3 else Path(sys.argv[1]).parent, sys.argv[2])
    print(json.dumps([str(r) for r in res], indent=2))
